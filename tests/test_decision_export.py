"""Take-home artefacts: decision PDF, Excel, and hashed share links.

The service layer of the questions is tested elsewhere. These lock the
properties that only exist at the boundary: the file is a real attachment,
the share token is hashed at rest, and a stranger with the URL sees the
frozen snapshot — not the household's ledger.
"""

from __future__ import annotations

import hashlib
import io
from datetime import timedelta
from urllib.parse import parse_qs, urlparse

import pytest
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APIClient

from apps.common.legal import DISCLAIMER
from apps.projections.share_models import DecisionShare

pytestmark = pytest.mark.django_db

BASE = "/api/v1/projections"
MORTGAGE = {
    "property_price_minor": 20_000_000,
    "deposit_minor": 5_000_000,
    "annual_rate": 0.09,
}


def _account(client, opening=8_000_000):
    return client.post(
        "/api/v1/finance/accounts/",
        {
            "name": "Checking",
            "account_type": "checking",
            "currency": "USD",
            "opening_balance_minor": opening,
        },
        format="json",
    ).data


def _token_from_url(url: str) -> str:
    return parse_qs(urlparse(url).query)["token"][0]


# ---------------------------------------------------------------------------
# PDF / Excel export
# ---------------------------------------------------------------------------


def test_decision_pdf_is_a_downloadable_document(tenant_context):
    _, client = tenant_context
    _account(client)
    res = client.post(f"{BASE}/questions/afford-mortgage/export.pdf", MORTGAGE, format="json")
    assert res.status_code == 200
    assert res["Content-Type"] == "application/pdf"
    assert "ledgerflow-afford-mortgage.pdf" in res["Content-Disposition"]
    assert res.content.startswith(b"%PDF")


def test_decision_xlsx_carries_the_disclaimer_and_the_headline(tenant_context):
    _, client = tenant_context
    _account(client)
    asked = client.post(f"{BASE}/questions/afford-mortgage/", MORTGAGE, format="json")
    res = client.post(f"{BASE}/questions/afford-mortgage/export.xlsx", MORTGAGE, format="json")
    assert res.status_code == 200
    assert "spreadsheetml" in res["Content-Type"]
    wb = load_workbook(io.BytesIO(res.content))
    summary = wb["Summary"]
    assert DISCLAIMER in (summary["A1"].value or "")
    assert asked.data["headline"] == summary["B5"].value
    assert asked.data["verdict"] == summary["B4"].value


def test_exporting_an_unknown_question_lists_the_ones_that_exist(tenant_context):
    _, client = tenant_context
    res = client.post(f"{BASE}/questions/should-i-buy-a-boat/export.pdf", {}, format="json")
    assert res.status_code == 404
    assert "afford-mortgage" in res.data["available"]


def test_exporting_on_an_empty_workspace_is_a_409(tenant_context):
    _, client = tenant_context
    res = client.post(f"{BASE}/questions/how-much-house/export.xlsx", {"annual_rate": 0.09}, format="json")
    assert res.status_code == 409


# ---------------------------------------------------------------------------
# Shareable snapshot
# ---------------------------------------------------------------------------


def test_share_returns_a_url_once_and_hashes_the_token(tenant_context, settings):
    _, client = tenant_context
    _account(client)
    settings.FRONTEND_BASE_URL = "https://app.example.com"
    res = client.post(f"{BASE}/questions/afford-mortgage/share/", MORTGAGE, format="json")
    assert res.status_code == 201
    url = res.data["url"]
    assert url.startswith("https://app.example.com/share/decision?token=")
    token = _token_from_url(url)
    digest = hashlib.sha256(token.encode()).hexdigest()
    stored = DecisionShare.objects.get()
    assert stored.token_hash == digest
    assert stored.token_hash != token
    assert not DecisionShare.objects.filter(token_hash=token).exists()


def test_a_stranger_can_open_the_frozen_snapshot(tenant_context):
    _, client = tenant_context
    _account(client)
    asked = client.post(f"{BASE}/questions/afford-mortgage/", MORTGAGE, format="json")
    shared = client.post(f"{BASE}/questions/afford-mortgage/share/", MORTGAGE, format="json")
    token = _token_from_url(shared.data["url"])

    public = APIClient()
    res = public.get(f"{BASE}/shared/{token}/")
    assert res.status_code == 200
    assert res.data["headline"] == asked.data["headline"]
    assert res.data["verdict"] == asked.data["verdict"]
    assert "email" not in res.data
    assert "tenant_id" not in res.data


def test_shared_exports_are_public_and_match_the_snapshot(tenant_context):
    _, client = tenant_context
    _account(client)
    shared = client.post(f"{BASE}/questions/afford-mortgage/share/", MORTGAGE, format="json")
    token = _token_from_url(shared.data["url"])
    public = APIClient()

    pdf = public.get(f"{BASE}/shared/{token}/export.pdf")
    assert pdf.status_code == 200
    assert pdf.content.startswith(b"%PDF")

    xlsx = public.get(f"{BASE}/shared/{token}/export.xlsx")
    assert xlsx.status_code == 200
    wb = load_workbook(io.BytesIO(xlsx.content))
    assert DISCLAIMER in (wb["Summary"]["A1"].value or "")


def test_an_expired_share_is_a_404(tenant_context):
    _, client = tenant_context
    _account(client)
    shared = client.post(f"{BASE}/questions/afford-mortgage/share/", MORTGAGE, format="json")
    token = _token_from_url(shared.data["url"])
    DecisionShare.objects.update(expires_at=timezone.now() - timedelta(seconds=1))
    assert APIClient().get(f"{BASE}/shared/{token}/").status_code == 404


def test_a_garbage_share_token_is_a_404():
    assert APIClient().get(f"{BASE}/shared/not-a-real-token/").status_code == 404


def test_creating_a_share_requires_authentication():
    assert APIClient().post(
        f"{BASE}/questions/afford-mortgage/share/", MORTGAGE, format="json"
    ).status_code in (
        401,
        403,
    )
