"""Take-home artefacts for a computed decision: PDF and Excel.

Both formats lead with the same disclaimer as the on-screen result. The
payload is the decision JSON (verdict, findings, assumptions), not a ledger
dump — these files are meant to be filed or shown to a lender, not to
reconstruct the household's books.
"""

from __future__ import annotations

import io
from xml.sax.saxutils import escape

from apps.common.legal import DISCLAIMER


def _money(amount_minor, currency: str) -> str:
    if amount_minor is None:
        return ""
    return f"{int(amount_minor) / 100:,.2f} {currency}"


def _findings(payload: dict, key: str) -> list[dict]:
    items = payload.get(key) or []
    return [item for item in items if isinstance(item, dict)]


def decision_pdf(payload: dict) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    buffer = io.BytesIO()
    question = str(payload.get("question") or "Decision")
    currency = str(payload.get("currency") or "")
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        title=question,
        leftMargin=18 * mm,
        rightMargin=18 * mm,
        topMargin=18 * mm,
        bottomMargin=18 * mm,
    )
    styles = getSampleStyleSheet()
    story = [
        Paragraph(escape(question), styles["Title"]),
        Spacer(1, 2 * mm),
        Paragraph(escape(DISCLAIMER), styles["BodyText"]),
        Spacer(1, 4 * mm),
    ]

    summary_rows = [
        ["Verdict", str(payload.get("verdict") or "")],
        ["Headline", str(payload.get("headline") or "")],
        ["Confidence", str(payload.get("confidence") or "")],
        ["Currency", currency],
    ]
    summary = Table(summary_rows, colWidths=[40 * mm, 130 * mm])
    summary.setStyle(
        TableStyle(
            [
                ("FONTSIZE", (0, 0), (-1, -1), 9),
                ("TEXTCOLOR", (0, 0), (0, -1), colors.HexColor("#656c81")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    story.append(summary)
    story.append(Spacer(1, 6 * mm))

    explanation = payload.get("explanation") or {}
    for paragraph in explanation.get("paragraphs") or []:
        story.append(Paragraph(escape(str(paragraph)), styles["BodyText"]))
        story.append(Spacer(1, 2 * mm))

    sections = (
        ("What it turns on", "because"),
        ("What it costs", "costs"),
        ("What could go wrong", "risks"),
        ("Worth considering instead", "alternatives"),
    )
    for title, key in sections:
        items = _findings(payload, key)
        if not items:
            continue
        story.append(Spacer(1, 3 * mm))
        story.append(Paragraph(escape(title), styles["Heading2"]))
        rows = [["Label", "Detail", "Amount"]]
        for item in items:
            rows.append(
                [
                    Paragraph(escape(str(item.get("label") or "")), styles["BodyText"]),
                    Paragraph(escape(str(item.get("text") or "")), styles["BodyText"]),
                    _money(item.get("amount_minor"), currency),
                ]
            )
        table = Table(rows, repeatRows=1, colWidths=[45 * mm, 85 * mm, 40 * mm])
        table.setStyle(
            TableStyle(
                [
                    ("FONTSIZE", (0, 0), (-1, -1), 8),
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f1f2f6")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("ALIGN", (2, 1), (2, -1), "RIGHT"),
                    ("TOPPADDING", (0, 0), (-1, -1), 3),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#fafbfc")]),
                ]
            )
        )
        story.append(table)

    assumptions = payload.get("assumptions") or []
    if assumptions:
        story.append(Spacer(1, 4 * mm))
        story.append(Paragraph("What this assumed", styles["Heading2"]))
        for assumption in assumptions:
            story.append(Paragraph(f"• {escape(str(assumption))}", styles["BodyText"]))

    doc.build(story)
    return buffer.getvalue()


def decision_xlsx(payload: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    summary["A1"] = DISCLAIMER
    summary.merge_cells("A1:B1")
    summary["A1"].alignment = Alignment(wrap_text=True, vertical="top")
    summary.row_dimensions[1].height = 64
    summary["A1"].font = Font(size=9, italic=True)

    rows = [
        (3, "Question", payload.get("question") or ""),
        (4, "Verdict", payload.get("verdict") or ""),
        (5, "Headline", payload.get("headline") or ""),
        (6, "Confidence", payload.get("confidence") or ""),
        (7, "Currency", payload.get("currency") or ""),
    ]
    for row, label, value in rows:
        summary[f"A{row}"] = label
        summary[f"B{row}"] = value
        summary[f"A{row}"].font = Font(color="656C81")
        summary[f"B{row}"].alignment = Alignment(wrap_text=True)

    explanation = payload.get("explanation") or {}
    prose = wb.create_sheet("Explanation")
    prose.append(["Paragraph"])
    for paragraph in explanation.get("paragraphs") or []:
        prose.append([str(paragraph)])

    findings = wb.create_sheet("Findings")
    findings.append(["Section", "Label", "Detail", "Amount", "Currency", "Months", "Percent"])
    currency = str(payload.get("currency") or "")
    for section, key in (
        ("because", "because"),
        ("costs", "costs"),
        ("risks", "risks"),
        ("alternatives", "alternatives"),
    ):
        for item in _findings(payload, key):
            amount = item.get("amount_minor")
            findings.append(
                [
                    section,
                    item.get("label") or "",
                    item.get("text") or "",
                    None if amount is None else int(amount) / 100,
                    currency if amount is not None else "",
                    item.get("months"),
                    item.get("percent"),
                ]
            )

    assumed = wb.create_sheet("Assumptions")
    assumed.append(["Assumption"])
    for assumption in payload.get("assumptions") or []:
        assumed.append([str(assumption)])

    for sheet in wb.worksheets:
        sheet.column_dimensions["A"].width = 22
        sheet.column_dimensions["B"].width = 48
        if sheet.title == "Findings":
            for index in range(1, 7):
                sheet.column_dimensions[get_column_letter(index)].width = 22
            sheet.column_dimensions["C"].width = 60

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
